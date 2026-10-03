from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

DEFAULT_OUTPUT = Path("artifacts/boolean_vaccine/boolean_vaccine_review.xlsx")


def read_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def add_table(ws, name: str, ref: str) -> None:
    table = Table(displayName=name, ref=ref)
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
    ws.add_table(table)


def write_sheet(ws, headers: list[str], rows: list[dict[str, Any]]) -> None:
    ws.append(headers)
    for row in rows:
        ws.append([row.get(h, "") for h in headers])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[column_cells[0].column_letter].width = min(max(max_len + 2, 12), 42)
    if rows:
        end_cell = ws.cell(row=len(rows) + 1, column=len(headers)).coordinate
        add_table(ws, f"{ws.title.replace(' ', '')}Table", f"A1:{end_cell}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build redacted Excel workbook for manual review of Boolean vaccine sprint outputs.")
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--grammar-summary", type=Path, default=Path("artifacts/boolean_vaccine/seed_grammar.redacted.summary.json"))
    parser.add_argument("--reward-summary", type=Path, default=Path("artifacts/boolean_vaccine/reward_loop_acceptance.redacted.summary.json"))
    parser.add_argument("--dod-summary", type=Path, default=Path("artifacts/boolean_vaccine/definition_of_done.summary.json"))
    parser.add_argument("--accepted", type=Path, default=Path("artifacts/boolean_vaccine/reward_loop_acceptance.redacted.jsonl"))
    parser.add_argument("--grammar", type=Path, default=Path("artifacts/boolean_vaccine/seed_grammar.redacted.jsonl"))
    args = parser.parse_args()
    grammar_summary = read_json(args.grammar_summary)
    reward_summary = read_json(args.reward_summary)
    dod_summary = read_json(args.dod_summary)
    accepted = load_jsonl(args.accepted)
    grammar_rows = load_jsonl(args.grammar)
    wb = Workbook()
    summary = wb.active
    summary.title = "Summary"
    summary_rows = [
        ["Metric", "Value"],
        ["Backend replay DoD", dod_summary.get("status")],
        ["Seed grammar status", grammar_summary.get("status")],
        ["Reward loop status", reward_summary.get("status")],
        ["Accepted total", reward_summary.get("accepted")],
        ["Target per family", reward_summary.get("target_per_family")],
        ["Min per family", reward_summary.get("min_per_family")],
        ["Max per family", reward_summary.get("max_per_family")],
        ["Public payloads redacted", reward_summary.get("payloads_redacted_in_public")],
        ["Claim boundary", reward_summary.get("claim_boundary")],
    ]
    for row in summary_rows:
        summary.append(row)
    for cell in summary[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    summary.column_dimensions["A"].width = 30
    summary.column_dimensions["B"].width = 95
    summary["B10"].alignment = Alignment(wrap_text=True)
    accepted_sheet = wb.create_sheet("Accepted_50")
    accepted_headers = ["sample_id", "seed_id", "family", "expected_truth", "observed_truth", "accepted", "reward_total", "waf_access_logged", "waf_audit_logged", "backend_reached", "oracle_matched", "payload_sha256", "token_class_signature", "operator_family", "length_bucket", "claim_scope"]
    write_sheet(accepted_sheet, accepted_headers, accepted)
    grammar_sheet = wb.create_sheet("Seed_Grammar")
    grammar_headers = ["seed_id", "family", "truth_label", "grammar_shape", "expr_sha256", "normalized_sha256", "grammar_valid", "operator_family", "token_count", "token_class_signature", "has_parentheses", "has_arithmetic", "has_range_predicate", "payload_redacted", "payload_exported"]
    write_sheet(grammar_sheet, grammar_headers, grammar_rows)
    checklist = wb.create_sheet("Checklist")
    checklist_rows = []
    for gate in dod_summary.get("gates", []):
        checklist_rows.append({"gate": gate.get("name"), "passed": gate.get("passed"), "evidence": json.dumps(gate.get("evidence", {}), sort_keys=True)[:32760]})
    write_sheet(checklist, ["gate", "passed", "evidence"], checklist_rows)
    checklist.column_dimensions["C"].width = 100
    for row in checklist.iter_rows(min_row=2, max_col=3):
        row[2].alignment = Alignment(wrap_text=True, vertical="top")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.output)
    print(json.dumps({"status": "ok", "output": str(args.output), "sheets": wb.sheetnames}, indent=2, sort_keys=True))


if __name__ == "__main__":
    main()
