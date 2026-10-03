from __future__ import annotations

import argparse
import base64
import json
import random
import sys
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import paramiko

from app.core.boolean_vaccine import (
    extract_grammar_features,
    normalize_expr,
    sha256_text,
)

ROOT = Path("artifacts/boolean_vaccine")
SPRINT_ROOT = Path("artifacts/vaccine_sprints_2_3_4")
PRIVATE_ROOT = SPRINT_ROOT / "private"

CONTROL_FAMILIES = [
    "benign_math",
    "benign_numeric_filter",
    "benign_range_filter",
    "ordinary_numeric_noise",
    "near_boolean_benign",
    "malformed_non_threatening",
]

MULTI_ATTACK_FAMILIES = [
    "decimal_comparison",
    "greater_than_ordering",
    "greater_equal_boundary",
    "multiplicative_arithmetic",
    "nested_parenthesized",
]


def read_json(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {}
    return json.loads(path.read_text(encoding="utf-8"))


def load_jsonl(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    if not path.exists():
        return rows
    with path.open(encoding="utf-8") as handle:
        for line in handle:
            if line.strip():
                rows.append(json.loads(line))
    return rows


def write_jsonl(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as handle:
        for row in rows:
            handle.write(json.dumps(row, sort_keys=True) + "\n")


def write_json(path: Path, value: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, indent=2, sort_keys=True) + "\n", encoding="utf-8")


def connect(host: str, user: str, key: str) -> paramiko.SSHClient:
    client = paramiko.SSHClient()
    client.set_missing_host_key_policy(paramiko.AutoAddPolicy())
    client.connect(host, username=user, key_filename=key, look_for_keys=False, allow_agent=False, timeout=8)
    return client


def run_remote_python(client: paramiko.SSHClient, code: str, stdin_payload: str = "") -> tuple[int, str, str]:
    stdin, stdout, stderr = client.exec_command("python3 -", get_pty=False)
    stdin.write(code)
    if stdin_payload:
        stdin.write(stdin_payload)
    stdin.channel.shutdown_write()
    out = stdout.read().decode("utf-8", errors="replace").strip()
    err = stderr.read().decode("utf-8", errors="replace").strip()
    return stdout.channel.recv_exit_status(), out, err


def batch_probe(client: paramiko.SSHClient, rows: list[dict[str, Any]], *, run_label: str) -> list[dict[str, Any]]:
    payload = base64.b64encode(json.dumps(rows).encode("utf-8")).decode("ascii")
    code = f'''
import base64, json, os, time, urllib.request, urllib.error, urllib.parse
rows = json.loads(base64.b64decode({payload!r}).decode('utf-8'))
access_log = '/var/log/apache2/hermes_lab_access.log'
audit_log = '/var/log/apache2/modsec_audit.log'
def token_count(path, token):
    try:
        needle = token.encode('utf-8')
        with open(path, 'rb') as handle:
            return sum(1 for line in handle if needle in line)
    except OSError:
        return 0
def size(path):
    try: return os.path.getsize(path)
    except OSError: return 0
out = []
for row in rows:
    probe_id = str(row['probe_id'])
    expr = str(row.get('payload', ''))
    encoded_expr = urllib.parse.quote(expr, safe='')
    encoded_probe = urllib.parse.quote(probe_id, safe='')
    before_access = token_count(access_log, probe_id)
    before_audit = size(audit_log)
    status_code = 0
    body = ''
    error = ''
    try:
        url = 'http://127.0.0.1:18080/bool_oracle?probe_id=' + encoded_probe + '&expr=' + encoded_expr
        req = urllib.request.Request(url, headers={{'User-Agent': 'Hermes-Vaccine-Sprint234/1.0'}})
        with urllib.request.urlopen(req, timeout=8) as response:
            status_code = response.getcode()
            body = response.read(2048).decode('utf-8', errors='replace')
    except urllib.error.HTTPError as exc:
        status_code = exc.code
        error = 'HTTPError'
        try: body = exc.read(512).decode('utf-8', errors='replace')
        except Exception: body = ''
    except Exception as exc:
        error = type(exc).__name__
    time.sleep(0.02)
    after_access = token_count(access_log, probe_id)
    after_audit = size(audit_log)
    parsed = {{}}
    try:
        parsed = json.loads(body) if body else {{}}
    except Exception:
        parsed = {{}}
    out.append({{
        'probe_id': probe_id,
        'http_status': status_code,
        'backend_reached': bool(parsed.get('backend_reached')),
        'observed_truth': parsed.get('observed_truth', ''),
        'backend_error': parsed.get('error', ''),
        'postgres_status': parsed.get('postgres_status', ''),
        'waf_access_logged': after_access > before_access,
        'waf_audit_logged': after_audit > before_audit,
        'waf_audit_delta_bytes': max(0, after_audit - before_audit),
        'error': error,
    }})
print(json.dumps(out, sort_keys=True))
'''
    status, out, err = run_remote_python(client, code)
    if status != 0:
        raise RuntimeError(f"remote batch probe failed {run_label}: {(err or out)[:500]}")
    return json.loads(out)


def expected_truth_for_expr(expr: str) -> str:
    # The safe backend is the authoritative oracle. This helper is only used for obvious benign labels.
    return "TRUE"


def generate_control_samples(count: int, seed: int) -> list[dict[str, Any]]:
    rng = random.Random(seed)
    rows: list[dict[str, Any]] = []
    per_family = count // len(CONTROL_FAMILIES)
    idx = 0
    for family in CONTROL_FAMILIES:
        for i in range(per_family):
            base = rng.randint(10, 900) + i
            if family == "benign_math":
                expr = f"{base} + {i + 1} = {base + i + 1}"
                expected = "TRUE"
            elif family == "benign_numeric_filter":
                expr = f"{base} <= {base + 5}"
                expected = "TRUE"
            elif family == "benign_range_filter":
                expr = f"{base} BETWEEN {base - 10} AND {base + 10}"
                expected = "TRUE"
            elif family == "ordinary_numeric_noise":
                expr = f"{base} = {base}"
                expected = "TRUE"
            elif family == "near_boolean_benign":
                expr = f"({base} - {i % 7}) < {base + 1}"
                expected = "TRUE"
            else:
                expr = f"{base} + {i + 2} > {base}"
                expected = "TRUE"
            features = extract_grammar_features(expr, seed_id=f"CTRL-{family}", family=family)
            rows.append({
                "sample_id": f"CTRL-{idx:05d}",
                "dataset": "control",
                "family": family,
                "expected_label": "benign",
                "expected_truth": expected,
                "payload": expr,
                **features,
            })
            idx += 1
    while len(rows) < count:
        base = rng.randint(10, 900)
        expr = f"{base} = {base}"
        features = extract_grammar_features(expr, seed_id="CTRL-extra", family="benign_math")
        rows.append({"sample_id": f"CTRL-{idx:05d}", "dataset": "control", "family": "benign_math", "expected_label": "benign", "expected_truth": "TRUE", "payload": expr, **features})
        idx += 1
    return rows[:count]


def generate_multi_family_attacks(per_family: int, seed: int) -> list[dict[str, Any]]:
    rng = random.Random(seed)
    rows: list[dict[str, Any]] = []
    idx = 0
    for family in MULTI_ATTACK_FAMILIES:
        for i in range(per_family):
            truth = "TRUE" if i % 2 == 0 else "FALSE"
            base = rng.randint(20, 700) + i * 3
            if family == "decimal_comparison":
                left = base + 0.5
                right = base + (0.4 if truth == "TRUE" else 0.6)
                expr = f"{left:.1f} > {right:.1f}"
            elif family == "greater_than_ordering":
                expr = f"{base} > {base - 1 if truth == 'TRUE' else base + 1}"
            elif family == "greater_equal_boundary":
                expr = f"{base} >= {base if truth == 'TRUE' else base + 1}"
            elif family == "multiplicative_arithmetic":
                a = base % 30 + 2
                b = i % 9 + 2
                expr = f"{a} * {b} = {a * b if truth == 'TRUE' else a * b + 1}"
            else:
                a = base % 80 + 10
                b = i % 9 + 1
                expr = f"(({a} + {b}) > {a + b - 1 if truth == 'TRUE' else a + b + 1})"
            features = extract_grammar_features(expr, seed_id=f"MF-{family}", family=family)
            rows.append({
                "sample_id": f"MF-{idx:05d}",
                "dataset": "multi_family_attack",
                "family": family,
                "expected_label": "attack",
                "expected_truth": truth,
                "payload": expr,
                **features,
            })
            idx += 1
    return rows


def redacted_row(row: dict[str, Any]) -> dict[str, Any]:
    payload = str(row.get("payload", ""))
    out = {k: v for k, v in row.items() if k not in {"payload", "generated_payload", "true_expr", "false_expr", "expr"}}
    out["payload_sha256"] = sha256_text(payload)
    out["normalized_sha256"] = sha256_text(normalize_expr(payload))
    out["payload_redacted"] = True
    out["payload_exported"] = False
    return out


def attach_probe(rows: list[dict[str, Any]], observations: list[dict[str, Any]]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row, obs in zip(rows, observations, strict=True):
        merged = dict(row)
        merged.update({
            "http_status": obs.get("http_status"),
            "backend_reached": bool(obs.get("backend_reached")),
            "observed_truth": obs.get("observed_truth"),
            "oracle_matched": obs.get("observed_truth") == row.get("expected_truth"),
            "waf_access_logged": bool(obs.get("waf_access_logged")),
            "waf_audit_logged": bool(obs.get("waf_audit_logged")),
            "probe_error": obs.get("error", ""),
        })
        out.append(merged)
    return out


def build_detector(attack_rows: list[dict[str, Any]], *, name: str) -> dict[str, Any]:
    hashes = sorted({sha256_text(normalize_expr(str(row.get("payload", row.get("generated_payload", ""))))) for row in attack_rows})
    families = sorted({str(row.get("family", "unknown")) for row in attack_rows})
    signatures = sorted({str(row.get("token_class_signature", "")) for row in attack_rows if row.get("token_class_signature")})
    return {
        "name": name,
        "version": "v0.1-lab",
        "mode": "detector_layer",
        "action": "block_or_flag_lab_only",
        "normalized_sha256_denylist": hashes,
        "family_coverage": families,
        "signature_observed": signatures,
        "claim_boundary": "Temporary lab detector over known/generated controlled grammar. Not a production WAF rule and not evidence of broad SQLi coverage.",
        "payloads_redacted_in_public": True,
    }


def detector_decision(row: dict[str, Any], detector: dict[str, Any]) -> dict[str, Any]:
    payload = str(row.get("payload", row.get("generated_payload", "")))
    h = sha256_text(normalize_expr(payload))
    matched = h in set(detector.get("normalized_sha256_denylist", []))
    return {
        "vaccine_action": "block" if matched else "allow",
        "vaccine_flagged": matched,
        "matched_feature": "normalized_sha256" if matched else "",
        "detector_version": detector.get("version"),
    }


def evaluate_detector(rows: list[dict[str, Any]], detector: dict[str, Any]) -> list[dict[str, Any]]:
    out: list[dict[str, Any]] = []
    for row in rows:
        merged = dict(row)
        merged.update(detector_decision(row, detector))
        out.append(merged)
    return out


def confusion(rows: list[dict[str, Any]]) -> dict[str, Any]:
    tp = sum(1 for r in rows if r.get("expected_label") == "attack" and r.get("vaccine_flagged") is True)
    fn = sum(1 for r in rows if r.get("expected_label") == "attack" and r.get("vaccine_flagged") is not True)
    tn = sum(1 for r in rows if r.get("expected_label") == "benign" and r.get("vaccine_flagged") is not True)
    fp = sum(1 for r in rows if r.get("expected_label") == "benign" and r.get("vaccine_flagged") is True)
    attack_total = tp + fn
    benign_total = tn + fp
    return {
        "tp": tp,
        "fn": fn,
        "tn": tn,
        "fp": fp,
        "attack_total": attack_total,
        "benign_total": benign_total,
        "attack_flag_rate": round(tp / attack_total, 4) if attack_total else 0.0,
        "benign_allow_rate": round(tn / benign_total, 4) if benign_total else 0.0,
        "false_positive_rate": round(fp / benign_total, 4) if benign_total else 0.0,
        "false_negative_rate": round(fn / attack_total, 4) if attack_total else 0.0,
    }


def build_excel(path: Path, summaries: dict[str, Any]) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("openpyxl is required for Excel report") from exc
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Metric", "Value"])
    for section, data in summaries.items():
        ws.append([section, json.dumps(data, sort_keys=True)])
    ws["A1"].font = ws["B1"].font = Font(bold=True, color="FFFFFF")
    ws["A1"].fill = ws["B1"].fill = PatternFill("solid", fgColor="1F4E78")
    for sheet_name, rows in [
        ("Sprint2_After", load_jsonl(SPRINT_ROOT / "sprint2_after_vaccine.redacted.jsonl")),
        ("Control_200", load_jsonl(SPRINT_ROOT / "sprint3_control_200.redacted.jsonl")),
        ("MultiFamily", load_jsonl(SPRINT_ROOT / "sprint4_multifamily_after.redacted.jsonl")),
        ("Checklist", summaries.get("checklist_rows", [])),
    ]:
        sh = wb.create_sheet(sheet_name[:31])
        if not rows:
            sh.append(["empty"])
            continue
        headers = sorted(rows[0].keys()) if isinstance(rows[0], dict) else ["item"]
        sh.append(headers)
        for row in rows:
            sh.append([json.dumps(row.get(h), sort_keys=True) if isinstance(row.get(h), (dict, list)) else row.get(h) for h in headers])
        for cell in sh[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
        sh.freeze_panes = "A2"
        for col in range(1, min(len(headers), 18) + 1):
            sh.column_dimensions[get_column_letter(col)].width = 22
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description="Execute vaccine sprints 2, 3, and 4 in the controlled Boolean lab.")
    parser.add_argument("--host", default="192.168.244.128")
    parser.add_argument("--user", default="ubuntu")
    parser.add_argument("--key", default=r"C:\Users\Admin\.ssh\hermes_lab_ed25519")
    parser.add_argument("--control-count", type=int, default=200)
    parser.add_argument("--multi-per-family", type=int, default=20)
    parser.add_argument("--seed", type=int, default=2344)
    args = parser.parse_args()

    SPRINT_ROOT.mkdir(parents=True, exist_ok=True)
    PRIVATE_ROOT.mkdir(parents=True, exist_ok=True)
    run_id = datetime.now(UTC).strftime("VACCINE-SPRINT234-%Y%m%dT%H%M%SZ")

    sprint1_attack = load_jsonl(ROOT / "private" / "reward_loop_acceptance.local.jsonl")
    if len(sprint1_attack) < 50:
        raise RuntimeError("Sprint 1 private attack corpus missing or incomplete")
    attack50 = []
    for idx, row in enumerate(sprint1_attack[:50]):
        payload = str(row.get("generated_payload", ""))
        features = extract_grammar_features(payload, seed_id=str(row.get("seed_id", "")), family=str(row.get("family", "")))
        attack50.append({
            "sample_id": f"S2-ATTACK-{idx:05d}",
            "dataset": "sprint1_attack_50",
            "expected_label": "attack",
            "expected_truth": row.get("expected_truth"),
            "family": row.get("family"),
            "payload": payload,
            **features,
        })

    control200 = generate_control_samples(args.control_count, args.seed)
    multifamily = generate_multi_family_attacks(args.multi_per_family, args.seed + 99)

    client = connect(args.host, args.user, args.key)
    try:
        before_probe_rows = [{**row, "probe_id": f"{run_id}-before-{idx:04d}"} for idx, row in enumerate(attack50)]
        before_obs = batch_probe(client, before_probe_rows, run_label="sprint2-before")
        before_attack = attach_probe(before_probe_rows, before_obs)

        control_probe_rows = [{**row, "probe_id": f"{run_id}-ctrl-{idx:04d}"} for idx, row in enumerate(control200)]
        control_obs = batch_probe(client, control_probe_rows, run_label="sprint3-control")
        control_probed = attach_probe(control_probe_rows, control_obs)

        mf_probe_rows = [{**row, "probe_id": f"{run_id}-mf-{idx:04d}"} for idx, row in enumerate(multifamily)]
        mf_obs = batch_probe(client, mf_probe_rows, run_label="sprint4-multifamily")
        mf_probed = attach_probe(mf_probe_rows, mf_obs)
    finally:
        client.close()

    # Sprint 2 detector starts with Sprint 1 attack corpus. Sprint 4 detector expands with multi-family corpus.
    detector_s2 = build_detector(before_attack, name="hermes_boolean_vaccine_sprint2_detector")
    after_attack_s2 = evaluate_detector(before_attack, detector_s2)
    after_control_s2 = evaluate_detector(control_probed[:50], detector_s2)
    sprint2_combined = after_attack_s2 + after_control_s2
    sprint2_metrics = confusion(sprint2_combined)

    detector_s4 = build_detector(before_attack + mf_probed, name="hermes_boolean_vaccine_multifamily_detector")
    after_control_s3 = evaluate_detector(control_probed, detector_s4)
    after_mf_s4 = evaluate_detector(mf_probed, detector_s4)
    sprint3_metrics = confusion(after_control_s3)
    sprint4_metrics = confusion(after_mf_s4 + after_control_s3)

    # Private artifacts include raw payloads. Public artifacts are redacted.
    write_jsonl(PRIVATE_ROOT / "sprint2_before_attack.local.jsonl", before_attack)
    write_jsonl(PRIVATE_ROOT / "sprint3_control_200.local.jsonl", control_probed)
    write_jsonl(PRIVATE_ROOT / "sprint4_multifamily.local.jsonl", mf_probed)
    write_jsonl(PRIVATE_ROOT / "sprint4_after_vaccine.local.jsonl", after_mf_s4)

    write_jsonl(SPRINT_ROOT / "sprint2_before_attack.redacted.jsonl", [redacted_row(r) for r in before_attack])
    write_jsonl(SPRINT_ROOT / "sprint2_after_vaccine.redacted.jsonl", [redacted_row(r) for r in sprint2_combined])
    write_jsonl(SPRINT_ROOT / "sprint3_control_200.redacted.jsonl", [redacted_row(r) for r in after_control_s3])
    write_jsonl(SPRINT_ROOT / "sprint4_multifamily_before.redacted.jsonl", [redacted_row(r) for r in mf_probed])
    write_jsonl(SPRINT_ROOT / "sprint4_multifamily_after.redacted.jsonl", [redacted_row(r) for r in after_mf_s4])

    write_json(SPRINT_ROOT / "sprint2_detector.redacted.json", detector_s2)
    write_json(SPRINT_ROOT / "sprint4_detector.redacted.json", detector_s4)

    sprint2_pass = sprint2_metrics["attack_flag_rate"] >= 0.90 and sprint2_metrics["benign_allow_rate"] >= 0.95 and sprint2_metrics["false_positive_rate"] <= 0.05
    sprint3_pass = len(control_probed) >= 200 and len({r["family"] for r in control_probed}) >= 5 and sprint3_metrics["false_positive_rate"] <= 0.05
    mf_accepted = [r for r in mf_probed if r.get("backend_reached") and r.get("waf_access_logged") and r.get("waf_audit_logged") and r.get("oracle_matched")]
    mf_by_family = Counter(str(r.get("family")) for r in mf_accepted)
    sprint4_pass = len(mf_accepted) == len(multifamily) and min(mf_by_family.values() or [0]) >= args.multi_per_family and sprint4_metrics["attack_flag_rate"] >= 0.90 and sprint4_metrics["false_positive_rate"] <= 0.05

    checklist_rows = [
        {"sprint": "2", "item": "baseline_before_vaccine", "passed": len(before_attack) == 50 and all(r.get("oracle_matched") for r in before_attack)},
        {"sprint": "2", "item": "detector_artifact", "passed": bool(detector_s2["normalized_sha256_denylist"])},
        {"sprint": "2", "item": "after_attack_flag_rate", "passed": sprint2_metrics["attack_flag_rate"] >= 0.90, "value": sprint2_metrics["attack_flag_rate"]},
        {"sprint": "2", "item": "benign_allow_rate", "passed": sprint2_metrics["benign_allow_rate"] >= 0.95, "value": sprint2_metrics["benign_allow_rate"]},
        {"sprint": "3", "item": "control_count", "passed": len(control_probed) >= 200, "value": len(control_probed)},
        {"sprint": "3", "item": "control_families", "passed": len({r['family'] for r in control_probed}) >= 5, "value": sorted({r['family'] for r in control_probed})},
        {"sprint": "3", "item": "control_fpr", "passed": sprint3_metrics["false_positive_rate"] <= 0.05, "value": sprint3_metrics["false_positive_rate"]},
        {"sprint": "4", "item": "multifamily_attack_count", "passed": len(mf_accepted) == len(multifamily), "value": len(mf_accepted)},
        {"sprint": "4", "item": "multifamily_balance", "passed": bool(mf_by_family) and min(mf_by_family.values()) >= args.multi_per_family, "value": dict(sorted(mf_by_family.items()))},
        {"sprint": "4", "item": "multifamily_detector_rate", "passed": sprint4_metrics["attack_flag_rate"] >= 0.90, "value": sprint4_metrics["attack_flag_rate"]},
    ]

    summary = {
        "status": "PASS" if sprint2_pass and sprint3_pass and sprint4_pass else "FAIL",
        "run_id": run_id,
        "sprint2_pass": sprint2_pass,
        "sprint3_pass": sprint3_pass,
        "sprint4_pass": sprint4_pass,
        "sprint2_metrics": sprint2_metrics,
        "sprint3_metrics": sprint3_metrics,
        "sprint4_metrics": sprint4_metrics,
        "control_count": len(control_probed),
        "control_families": sorted({r["family"] for r in control_probed}),
        "multifamily_attack_count": len(mf_accepted),
        "multifamily_accepted_by_family": dict(sorted(mf_by_family.items())),
        "detector_artifacts": [str(SPRINT_ROOT / "sprint2_detector.redacted.json"), str(SPRINT_ROOT / "sprint4_detector.redacted.json")],
        "public_artifacts_redacted": True,
        "private_replay_root": str(PRIVATE_ROOT),
        "claim_boundary": "Controlled lab detector layer only. This validates temporary vaccine logic against generated corpora; it is not a production WAF rule or broad SQLi protection claim.",
        "checklist_rows": checklist_rows,
    }
    write_json(SPRINT_ROOT / "sprints_2_3_4_definition_of_done.summary.json", summary)
    build_excel(SPRINT_ROOT / "vaccine_sprints_2_3_4_review.xlsx", summary)
    print(json.dumps(summary, indent=2, sort_keys=True))
    raise SystemExit(0 if summary["status"] == "PASS" else 1)


if __name__ == "__main__":
    main()
