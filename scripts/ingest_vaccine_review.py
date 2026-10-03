from __future__ import annotations

import argparse
import json
from collections import Counter
from datetime import UTC, datetime
from pathlib import Path
from typing import Any


def load_rows_from_workbook(path: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise SystemExit("openpyxl is required: uv run --with openpyxl python scripts/ingest_vaccine_review.py") from exc
    wb = load_workbook(path)
    if "Candidate_Review" not in wb.sheetnames:
        raise ValueError("workbook missing Candidate_Review sheet")
    ws = wb["Candidate_Review"]
    headers = [str(cell.value) if cell.value is not None else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows: list[dict[str, Any]] = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        if not any(value is not None for value in values):
            continue
        rows.append({headers[idx]: values[idx] for idx in range(min(len(headers), len(values)))})
    return rows


def metrics(rows: list[dict[str, Any]]) -> dict[str, Any]:
    tp = tn = fp = fn = undecided = 0
    for row in rows:
        expected = str(row.get("expected_label") or "").lower()
        verdict = str(row.get("manual_verdict") or "").lower()
        if verdict not in {"approve", "reject"}:
            undecided += 1
            continue
        flagged = verdict == "reject"
        if expected == "attack" and flagged:
            tp += 1
        elif expected == "attack" and not flagged:
            fn += 1
        elif expected == "benign" and flagged:
            fp += 1
        elif expected == "benign" and not flagged:
            tn += 1
    attack_total = tp + fn
    benign_total = tn + fp
    return {
        "tp": tp,
        "tn": tn,
        "fp": fp,
        "fn": fn,
        "undecided": undecided,
        "attack_total": attack_total,
        "benign_total": benign_total,
        "attack_flag_rate": round(tp / attack_total, 4) if attack_total else 0.0,
        "benign_allow_rate": round(tn / benign_total, 4) if benign_total else 0.0,
        "false_positive_rate": round(fp / benign_total, 4) if benign_total else 0.0,
        "false_negative_rate": round(fn / attack_total, 4) if attack_total else 0.0,
    }


def main() -> None:
    parser = argparse.ArgumentParser(description="Ingest a Hermes vaccine human-review workbook.")
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--output", type=Path, default=Path("artifacts/vaccine_sprints_8_9_10/review_ingest.summary.json"))
    parser.add_argument("--audit", type=Path, default=Path("artifacts/vaccine_sprints_8_9_10/review_audit.redacted.jsonl"))
    args = parser.parse_args()
    rows = load_rows_from_workbook(args.workbook)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    audit_rows: list[dict[str, Any]] = []
    family_counter: Counter[str] = Counter()
    for row in rows:
        family = str(row.get("family") or "unknown")
        family_counter[family] += 1
        audit_rows.append({
            "sample_id": row.get("sample_id"),
            "payload_sha256": row.get("payload_sha256"),
            "family": family,
            "expected_label": row.get("expected_label"),
            "manual_verdict": row.get("manual_verdict"),
            "accepted_for_vaccine": row.get("accepted_for_vaccine"),
            "reviewer_notes_present": bool(row.get("reviewer_notes")),
            "payload_redacted": True,
        })
    with args.audit.open("w", encoding="utf-8") as handle:
        for row in audit_rows:
            handle.write(json.dumps(row, sort_keys=True) + "\n")
    summary = {
        "status": "PASS" if rows else "FAIL",
        "review_rows": len(rows),
        "family_counts": dict(sorted(family_counter.items())),
        "metrics_after_review": metrics(rows),
        "audit_output": str(args.audit),
        "payloads_redacted_in_public": True,
        "claim_boundary": "Workflow validation only; manual verdicts must be reviewed by the user before operational use.",
        "ingested_at": datetime.now(UTC).isoformat(),
    }
    args.output.write_text(json.dumps(summary, indent=2, sort_keys=True) + "\n", encoding="utf-8")
    print(json.dumps(summary, indent=2, sort_keys=True))
    raise SystemExit(0 if summary["status"] == "PASS" else 1)


if __name__ == "__main__":
    main()
